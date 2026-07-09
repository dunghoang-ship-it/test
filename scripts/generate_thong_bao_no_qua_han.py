#!/usr/bin/env python3
"""Generate BIDV 'Thông báo nợ quá hạn' DOCX from one overdue loan row."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


@dataclass
class LoanNoticeData:
    so_van_ban: str
    ten_khach_hang: str
    ngay_bao_cao: date
    so_khoan_vay: str
    du_no_goc: int
    du_no_lai_va_phat: int
    qua_han_goc: int
    qua_han_lai_va_phat: int
    so_ngay_qua_han: int
    so_dien_thoai: str = "02513 526 325"
    nguoi_ky: str = "Nguyễn Quốc Tuấn"
    chuc_danh_ky: str = "P.KHCN"
    ngay_them_han_thanh_toan: int = 15


def format_vnd(n: int) -> str:
    return f"{n:,}".replace(",", ".")


ONES = [
    "",
    "một",
    "hai",
    "ba",
    "bốn",
    "năm",
    "sáu",
    "bảy",
    "tám",
    "chín",
]


def _read_two_digits(n: int, *, use_le: bool) -> str:
    if n < 10:
        if not n:
            return ""
        return f"lẻ {ONES[n]}" if use_le else ONES[n]
    if n < 20:
        return "mười" if n == 10 else f"mười {ONES[n % 10] if n % 10 != 5 else 'lăm'}"
    tens, ones = divmod(n, 10)
    parts = [ONES[tens], "mươi"]
    if ones == 1:
        parts.append("mốt")
    elif ones == 4:
        parts.append("tư")
    elif ones == 5:
        parts.append("lăm")
    elif ones:
        parts.append(ONES[ones])
    return " ".join(parts)


def _read_three_digits(n: int) -> str:
    hundreds, rem = divmod(n, 100)
    parts: list[str] = []
    if hundreds:
        parts.append(f"{ONES[hundreds]} trăm")
    if rem:
        # "lẻ" only after hundreds, e.g. "một trăm lẻ năm"
        parts.append(_read_two_digits(rem, use_le=bool(hundreds)))
    return " ".join(parts)


def number_to_vietnamese_words(n: int) -> str:
    if n == 0:
        return "Không đồng"

    scales = [
        (1_000_000_000_000, "nghìn tỷ"),
        (1_000_000_000, "tỷ"),
        (1_000_000, "triệu"),
        (1_000, "nghìn"),
    ]
    parts: list[str] = []
    remaining = n
    for value, label in scales:
        if remaining >= value:
            chunk = remaining // value
            remaining %= value
            chunk_words = _read_three_digits(chunk)
            parts.append(f"{chunk_words} {label}")
    if remaining:
        parts.append(_read_three_digits(remaining))

    text = " ".join(parts).strip()
    text = " ".join(text.split())
    return text[:1].upper() + text[1:] + " đồng"


def set_run_font(run, *, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(p, *, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0, line=1.15, first_line=0):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if first_line:
        pf.first_line_indent = Cm(first_line)


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def set_cell_border(cell, **borders):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, attrs in borders.items():
        element = OxmlElement(f"w:{edge}")
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def shade_cell(cell, fill="D9E2F3"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_format(p, align=align, space_before=2, space_after=2, line=1.0)
    add_text(p, text, size=size, bold=bold)


def merge_and_fill(table, r1, c1, r2, c2, text, **kwargs):
    cell = table.cell(r1, c1)
    cell.merge(table.cell(r2, c2))
    set_cell_text(cell, text, **kwargs)
    return cell


def build_document(data: LoanNoticeData, output_path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    # Header table: bank info | national motto
    header = doc.add_table(rows=1, cols=2)
    header.autofit = True
    left, right = header.rows[0].cells

    left.text = ""
    p = left.paragraphs[0]
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15)
    add_text(p, "NGÂN HÀNG TMCP ĐẦU TƯ\n", bold=True, size=11)
    add_text(p, "VÀ PHÁT TRIỂN VIỆT NAM\n", bold=True, size=11)
    add_text(p, "CHI NHÁNH ĐÔNG ĐỒNG NAI", bold=True, size=11)

    p2 = left.add_paragraph()
    set_paragraph_format(p2, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, line=1.15)
    add_text(p2, f"Số: {data.so_van_ban}", size=11)
    p3 = left.add_paragraph()
    set_paragraph_format(p3, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1)
    add_text(
        p3,
        f"(V/v Thông báo nợ quá hạn đối với khách hàng ông/bà {data.ten_khach_hang})",
        size=10,
        italic=True,
    )

    right.text = ""
    p = right.paragraphs[0]
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15)
    add_text(p, "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM\n", bold=True, size=11)
    add_text(p, "Độc lập - Tự do - Hạnh phúc", bold=True, size=11)

    # underline under motto
    p_line = right.add_paragraph()
    set_paragraph_format(p_line, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    add_text(p_line, "________________________", size=10)

    p_date = right.add_paragraph()
    set_paragraph_format(p_date, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, line=1.15)
    add_text(
        p_date,
        f"Đồng Nai, ngày {data.ngay_bao_cao.day:02d} tháng {data.ngay_bao_cao.month:02d} năm {data.ngay_bao_cao.year}",
        size=11,
        italic=True,
    )

    # Title
    title = doc.add_paragraph()
    set_paragraph_format(title, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=6, line=1.15)
    add_text(title, "THÔNG BÁO NỢ QUÁ HẠN", bold=True, size=14)

    greeting = doc.add_paragraph()
    set_paragraph_format(greeting, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10, line=1.15)
    add_text(greeting, f"Kính gửi: ông/bà {data.ten_khach_hang}", bold=True, size=12)

    body1 = doc.add_paragraph()
    set_paragraph_format(body1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=1.0, line=1.15)
    add_text(
        body1,
        "Căn cứ Hợp đồng tín dụng ký giữa Khách hàng và Ngân hàng TMCP Đầu tư và Phát triển Việt Nam;",
        size=12,
    )

    ngay_str = data.ngay_bao_cao.strftime("%d/%m/%Y")
    body2 = doc.add_paragraph()
    set_paragraph_format(body2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=1.0, space_after=8, line=1.15)
    add_text(
        body2,
        "Ngân hàng TMCP Đầu tư và Phát triển Việt Nam – Chi nhánh Đông Đồng Nai "
        "(sau đây gọi là “Ngân hàng”) trân trọng thông báo đến quý khách được biết về khoản nợ "
        f"tính đến ngày {ngay_str} và đề nghị như sau:",
        size=12,
    )

    unit = doc.add_paragraph()
    set_paragraph_format(unit, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4)
    add_text(unit, "ĐVT: Đồng", italic=True, size=10)

    # Debt table: 8 columns visually grouped
    # Row0 headers: Số khoản vay | Dư nợ hiện tại (3) | Dư nợ quá hạn (3) | Số ngày quá hạn
    table = doc.add_table(rows=4, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    thin = {"sz": "4", "val": "single", "color": "000000"}

    # Header row 0
    merge_and_fill(table, 0, 0, 1, 0, "Số khoản vay", bold=True, size=9)
    merge_and_fill(table, 0, 1, 0, 3, "Dư nợ hiện tại", bold=True, size=9)
    merge_and_fill(table, 0, 4, 0, 6, "Dư nợ quá hạn", bold=True, size=9)
    merge_and_fill(table, 0, 7, 1, 7, "Số ngày\nquá hạn", bold=True, size=9)

    for idx, label in enumerate(["Gốc", "Lãi và\nlãi phạt", "Tổng"], start=1):
        set_cell_text(table.cell(1, idx), label, bold=True, size=9)
    for idx, label in enumerate(["Gốc", "Lãi và\nlãi phạt", "Tổng"], start=4):
        set_cell_text(table.cell(1, idx), label, bold=True, size=9)

    for r in range(2):
        for c in range(8):
            shade_cell(table.cell(r, c), "D9E2F3")

    tong_hien_tai = data.du_no_goc + data.du_no_lai_va_phat
    tong_qua_han = data.qua_han_goc + data.qua_han_lai_va_phat

    values_row = [
        data.so_khoan_vay,
        format_vnd(data.du_no_goc),
        format_vnd(data.du_no_lai_va_phat),
        format_vnd(tong_hien_tai),
        format_vnd(data.qua_han_goc),
        format_vnd(data.qua_han_lai_va_phat),
        format_vnd(tong_qua_han),
        str(data.so_ngay_qua_han),
    ]
    for c, val in enumerate(values_row):
        set_cell_text(table.cell(2, c), val, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT if c else WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(2, 0), data.so_khoan_vay, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(2, 7), str(data.so_ngay_qua_han), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    totals = [
        "Tổng cộng",
        format_vnd(data.du_no_goc),
        format_vnd(data.du_no_lai_va_phat),
        format_vnd(tong_hien_tai),
        format_vnd(data.qua_han_goc),
        format_vnd(data.qua_han_lai_va_phat),
        format_vnd(tong_qua_han),
        "",
    ]
    for c, val in enumerate(totals):
        align = WD_ALIGN_PARAGRAPH.CENTER if c in (0, 7) else WD_ALIGN_PARAGRAPH.RIGHT
        set_cell_text(table.cell(3, c), val, bold=True, size=9, align=align)

    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=thin, bottom=thin, left=thin, right=thin)

    han_thanh_toan = data.ngay_bao_cao + timedelta(days=data.ngay_them_han_thanh_toan)
    han_str = han_thanh_toan.strftime("%d/%m/%Y")

    summary = doc.add_paragraph()
    set_paragraph_format(summary, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=1.0, space_before=12, line=1.15)
    add_text(summary, f"Tổng nợ gốc, lãi của quý khách đến ngày {ngay_str} là ", size=12)
    add_text(summary, f"{format_vnd(tong_hien_tai)} đồng", bold=True, size=12)
    add_text(
        summary,
        f" (Bằng chữ: {number_to_vietnamese_words(tong_hien_tai)}). Trong đó, tổng số tiền quá hạn là ",
        size=12,
    )
    add_text(summary, f"{format_vnd(tong_qua_han)} đồng", bold=True, size=12)
    add_text(
        summary,
        f" (Bằng chữ: {number_to_vietnamese_words(tong_qua_han)}), đề nghị quý khách thu xếp để thanh toán "
        f"toàn bộ số tiền nợ quá hạn tại Ngân hàng trước ngày {han_str}. Trường hợp quý khách không thanh toán "
        "đầy đủ nợ quá hạn trong thời gian thông báo này, Ngân hàng sẽ thực hiện các biện pháp khởi kiện, "
        "phát mại tài sản bảo đảm để thu hồi nợ.",
        size=12,
    )

    contact = doc.add_paragraph()
    set_paragraph_format(contact, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=1.0, space_before=8, line=1.15)
    add_text(
        contact,
        "Mọi thông tin xin vui lòng liên hệ đến Phòng Khách hàng cá nhân – Ngân hàng TMCP Đầu tư và Phát triển "
        f"Việt Nam – Chi nhánh Đông Đồng Nai, số điện thoại: {data.so_dien_thoai}.",
        size=12,
    )

    closing = doc.add_paragraph()
    set_paragraph_format(closing, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=10)
    add_text(closing, "Trân trọng thông báo./.", size=12)

    # Footer: nơi nhận | chữ ký
    footer = doc.add_table(rows=1, cols=2)
    left_f, right_f = footer.rows[0].cells

    left_f.text = ""
    p = left_f.paragraphs[0]
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15)
    add_text(p, "Nơi nhận: (…bản)", bold=True, italic=True, size=11)
    for line in [
        "- Như trên;",
        "- Ban Giám đốc B.One",
        "- P. KHCN B.One;",
        "- Lưu VT; P. QTTD.",
    ]:
        lp = left_f.add_paragraph()
        set_paragraph_format(lp, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.1)
        add_text(lp, line, size=11)

    right_f.text = ""
    p = right_f.paragraphs[0]
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15)
    add_text(p, "ĐẠI DIỆN NGÂN HÀNG", bold=True, size=12)

    for _ in range(5):
        sp = right_f.add_paragraph()
        set_paragraph_format(sp, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(sp, "", size=12)

    name_p = right_f.add_paragraph()
    set_paragraph_format(name_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6)
    add_text(name_p, data.nguoi_ky, bold=True, size=12)

    role_p = right_f.add_paragraph()
    set_paragraph_format(role_p, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(role_p, data.chuc_danh_ky, size=11)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def main():
    # Dòng dữ liệu quá hạn từ sheet gid=1981502462:
    # NGUYEN THI THANH THUY | TK 406004006532 | trạng thái nợ B | quá hạn 1 ngày
    data = LoanNoticeData(
        so_van_ban="999/TB-BIDV.ĐĐN",
        ten_khach_hang="Nguyễn Thị Thanh Thúy",
        ngay_bao_cao=date(2026, 7, 7),  # lịch trả 06/07/2026, quá hạn 1 ngày
        so_khoan_vay="406004006532",
        du_no_goc=898_494_837,
        # Số tiền lãi hiện tại (làm tròn) + dư lãi phạt cho gốc
        du_no_lai_va_phat=1_952_580 + 88_619,
        qua_han_goc=898_494_837,
        # Số tiền lãi quá hạn + dư lãi phạt cho gốc
        qua_han_lai_va_phat=1_775_342 + 88_619,
        so_ngay_qua_han=1,
        ngay_them_han_thanh_toan=15,
    )

    out = Path("/workspace/output/Thong_bao_no_qua_han_Nguyen_Thi_Thanh_Thuy.docx")
    path = build_document(data, out)

    tong_hien_tai = data.du_no_goc + data.du_no_lai_va_phat
    tong_qua_han = data.qua_han_goc + data.qua_han_lai_va_phat
    print(f"Created: {path}")
    print(f"Khách hàng: {data.ten_khach_hang}")
    print(f"Số khoản vay: {data.so_khoan_vay}")
    print(f"Tổng dư nợ hiện tại: {format_vnd(tong_hien_tai)}")
    print(f"Tổng quá hạn: {format_vnd(tong_qua_han)}")
    print(f"Bằng chữ tổng nợ: {number_to_vietnamese_words(tong_hien_tai)}")
    print(f"Bằng chữ quá hạn: {number_to_vietnamese_words(tong_qua_han)}")


if __name__ == "__main__":
    main()
